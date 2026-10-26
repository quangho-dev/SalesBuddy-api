from datetime import datetime, timezone
from io import BytesIO
import math
import re
import unicodedata

from openpyxl import Workbook, load_workbook
from src.services.supabase import supabase
from src.services.users import ensure_user

TABLE = "products"
TEXT_FIELDS = (
    "product_code",
    "name",
    "product_type",
    "unit",
    "notes",
)
NUMBER_FIELDS = (
    "selling_price",
    "stock",
)


class ProductNotFoundError(LookupError):
    pass


class ProductImportError(ValueError):
    pass


TEMPLATE_HEADERS = (
    "Mã sản phẩm",
    "Tên sản phẩm",
    "Đơn vị tính",
    "Giá bán",
    "Tồn kho",
    "Ghi chú",
    "Loại sản phẩm",
)

HEADER_ALIASES = {
    "product_code": {
        "ma san pham",
        "product code",
        "product_code",
        "sku",
        "code",
        "ma",
    },
    "name": {
        "ten san pham",
        "name",
        "ten",
    },
    "product_type": {
        "loai san pham",
        "product type",
        "product_type",
        "type",
        "category",
        "loai",
    },
    "unit": {
        "don vi tinh",
        "on vi tinh",
        "unit",
        "dvt",
        "don vi",
    },
    "selling_price": {
        "gia ban",
        "selling price",
        "selling_price",
        "price",
        "gia",
    },
    "stock": {
        "ton kho",
        "stock",
        "inventory",
        "sl ton",
    },
    "notes": {
        "ghi chu",
        "notes",
        "note",
    },
}

INSERT_CHUNK_SIZE = 100


def _now() -> str:
    return datetime.now(timezone.utc).isoformat()


def _text(value) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _number(value) -> float:
    if value is None or value == "":
        return 0
    return float(value)


def _to_row(payload: dict, clerk_id: str) -> dict:
    row = {field: _text(payload.get(field)) for field in TEXT_FIELDS}
    row.update({field: _number(payload.get(field)) for field in NUMBER_FIELDS})
    row["clerk_id"] = clerk_id
    return row


def _owned_query(product_id: str, clerk_id: str):
    return (
        supabase.table(TABLE)
        .select("*")
        .eq("id", product_id)
        .eq("clerk_id", clerk_id)
        .limit(1)
    )


def _fold_header(value: str) -> str:
    # Đ/đ does not decompose in NFD, so it would be stripped from a-z.
    text = value.strip().lower().replace("đ", "d")
    normalized = unicodedata.normalize("NFD", text)
    stripped = "".join(ch for ch in normalized if unicodedata.category(ch) != "Mn")
    return re.sub(r"[^a-z0-9]+", " ", stripped).strip()


def _parse_number(value) -> float:
    if value is None or value == "":
        return 0
    if isinstance(value, bool):
        return 0
    if isinstance(value, (int, float)):
        if isinstance(value, float) and math.isnan(value):
            return 0
        return float(value)
    text = (
        str(value)
        .strip()
        .replace("đ", "")
        .replace("₫", "")
        .replace(" ", "")
        .replace("\u00a0", "")
    )
    if not text:
        return 0
    if "," in text and "." in text:
        if text.rfind(",") > text.rfind("."):
            text = text.replace(".", "").replace(",", ".")
        else:
            text = text.replace(",", "")
    elif text.count(".") > 1:
        text = text.replace(".", "")
    elif "," in text:
        if text.count(",") > 1:
            text = text.replace(",", "")
        else:
            text = text.replace(",", ".")
    return float(text)


def _cell_text(value) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _map_headers(headers: tuple) -> dict[int, str]:
    mapping = {}
    for index, header in enumerate(headers):
        folded = _fold_header(_cell_text(header))
        if not folded:
            continue
        for field, aliases in HEADER_ALIASES.items():
            if folded in aliases:
                mapping[index] = field
                break
    return mapping


def _row_payload(values: tuple, mapping: dict[int, str]) -> dict:
    payload = {field: "" if field in TEXT_FIELDS else 0 for field in (*TEXT_FIELDS, *NUMBER_FIELDS)}
    for index, field in mapping.items():
        if index >= len(values):
            continue
        value = values[index]
        if field in NUMBER_FIELDS:
            payload[field] = _parse_number(value)
        else:
            payload[field] = _cell_text(value)
    return payload


def _is_empty_payload(payload: dict) -> bool:
    has_text = any(_text(payload.get(field)) for field in TEXT_FIELDS)
    has_number = any(_number(payload.get(field)) for field in NUMBER_FIELDS)
    return not has_text and not has_number


def _identity_key(product_code, name) -> tuple[str, str] | None:
    code = re.sub(r"\s+", " ", _text(product_code)).casefold()
    product_name = re.sub(r"\s+", " ", _text(name)).casefold()
    if not code and not product_name:
        return None
    return (code, product_name)


def _existing_product_keys(clerk_id: str) -> set[tuple[str, str]]:
    keys: set[tuple[str, str]] = set()
    page_size = 1000
    start = 0
    while True:
        result = (
            supabase.table(TABLE)
            .select("product_code, name")
            .eq("clerk_id", clerk_id)
            .range(start, start + page_size - 1)
            .execute()
        )
        rows = result.data or []
        for row in rows:
            key = _identity_key(row.get("product_code"), row.get("name"))
            if key:
                keys.add(key)
        if len(rows) < page_size:
            break
        start += page_size
    return keys


def create_products(clerk_id: str, payloads: list[dict]) -> list[dict]:
    if not payloads:
        return []
    ensure_user(clerk_id)
    saved = []
    for start in range(0, len(payloads), INSERT_CHUNK_SIZE):
        chunk = [_to_row(item, clerk_id) for item in payloads[start : start + INSERT_CHUNK_SIZE]]
        result = supabase.table(TABLE).insert(chunk).execute()
        saved.extend(result.data or [])
    return saved


def import_products_from_xlsx(clerk_id: str, content: bytes) -> dict:
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise ProductImportError("File Excel không hợp lệ.") from exc

    try:
        sheet = workbook.active
        rows = sheet.iter_rows(values_only=True)
        try:
            headers = next(rows)
        except StopIteration as exc:
            raise ProductImportError("File Excel trống.") from exc

        mapping = _map_headers(headers or ())
        if not mapping:
            raise ProductImportError(
                "Không nhận diện được cột. Hãy dùng file mẫu (Mã sản phẩm, Tên sản phẩm, ...)."
            )

        payloads = []
        errors = []
        skipped = 0
        duplicates = 0
        seen = _existing_product_keys(clerk_id)
        for row_number, values in enumerate(rows, start=2):
            values = values or ()
            if all(value is None or str(value).strip() == "" for value in values):
                skipped += 1
                continue
            try:
                payload = _row_payload(values, mapping)
            except (TypeError, ValueError):
                errors.append({"row": row_number, "message": "Số không hợp lệ."})
                continue
            if _is_empty_payload(payload):
                skipped += 1
                continue
            key = _identity_key(payload.get("product_code"), payload.get("name"))
            if key and key in seen:
                duplicates += 1
                continue
            if key:
                seen.add(key)
            payloads.append(payload)
    finally:
        workbook.close()

    saved = create_products(clerk_id, payloads)
    return {
        "imported": len(saved),
        "skipped": skipped,
        "duplicates": duplicates,
        "errors": errors,
        "products": saved,
    }


def build_product_template() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "San pham"
    sheet.append(list(TEMPLATE_HEADERS))
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def create_product(clerk_id: str, payload: dict) -> dict:
    ensure_user(clerk_id)
    result = supabase.table(TABLE).insert(_to_row(payload, clerk_id)).execute()
    if not result.data:
        raise RuntimeError("Failed to create product")
    return result.data[0]


def list_products(clerk_id: str) -> list[dict]:
    result = (
        supabase.table(TABLE)
        .select("*")
        .eq("clerk_id", clerk_id)
        .order("created_at", desc=True)
        .execute()
    )
    return result.data or []


def get_product(product_id: str, clerk_id: str) -> dict:
    result = _owned_query(product_id, clerk_id).execute()
    data = result.data or []
    if not data:
        raise ProductNotFoundError("Product not found")
    return data[0]


def update_product(product_id: str, clerk_id: str, payload: dict) -> dict:
    get_product(product_id, clerk_id)

    updates = {}
    for field in TEXT_FIELDS:
        if field in payload and payload[field] is not None:
            updates[field] = _text(payload[field])
    for field in NUMBER_FIELDS:
        if field in payload and payload[field] is not None:
            updates[field] = _number(payload[field])
    if not updates:
        return get_product(product_id, clerk_id)

    updates["updated_at"] = _now()
    result = (
        supabase.table(TABLE)
        .update(updates)
        .eq("id", product_id)
        .eq("clerk_id", clerk_id)
        .execute()
    )
    if not result.data:
        raise ProductNotFoundError("Product not found")
    return result.data[0]


def delete_product(product_id: str, clerk_id: str) -> dict:
    existing = get_product(product_id, clerk_id)
    supabase.table(TABLE).delete().eq("id", product_id).eq(
        "clerk_id", clerk_id
    ).execute()
    return existing
