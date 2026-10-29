from datetime import datetime, timezone
from io import BytesIO
from uuid import uuid4
from zoneinfo import ZoneInfo
import re
import unicodedata

from openpyxl import Workbook, load_workbook
from src.services.settings import get_settings
from src.services.supabase import supabase
from src.services.users import ensure_user

TABLE = "customers"
TEXT_FIELDS = (
    "company_name",
    "contact_person",
    "address",
    "phone_number",
    "tax_code",
    "website",
    "facebook",
    "zalo",
    "notes",
)
INT_FIELDS = ("payment_term_days",)
FIELDS = TEXT_FIELDS + INT_FIELDS


class CustomerNotFoundError(LookupError):
    pass


class CustomerImportError(ValueError):
    pass


TEMPLATE_HEADERS = (
    "Tên công ty",
    "Điện thoại",
    "Người liên hệ",
    "Địa chỉ",
    "Mã số thuế",
    "Hạn công nợ (ngày)",
    "Lịch sử tương tác",
)

_LOCAL_TZ = ZoneInfo("Asia/Ho_Chi_Minh")
_INTERACTION_LINE = re.compile(
    r"^(\d{1,2}[/-]\d{1,2}[/-]\d{4}|\d{4}-\d{1,2}-\d{1,2})"
    r"(?:[ T](\d{1,2}:\d{2}))?"
    r"\s*(.*)$"
)

HEADER_ALIASES = {
    "company_name": {
        "ten cong ty",
        "company name",
        "company",
        "cong ty",
    },
    "phone_number": {
        "dien thoai",
        "so dien thoai",
        "phone",
        "phone number",
        "sdt",
        "tel",
        "mobile",
    },
    "contact_person": {
        "nguoi lien he",
        "nguoi lien lac",
        "contact",
        "contact person",
        "lien he",
    },
    "address": {
        "dia chi",
        "address",
        "dia chi cong ty",
    },
    "tax_code": {
        "ma so thue",
        "mst",
        "tax code",
        "tax id",
        "vat",
    },
    "payment_term_days": {
        "han cong no",
        "han cong no ngay",
        "so ngay no",
        "payment term",
        "credit days",
        "debt term",
    },
    "interaction_history": {
        "lich su tuong tac",
        "tuong tac",
        "interaction",
        "interaction history",
        "history",
    },
}

INSERT_CHUNK_SIZE = 100


def _now() -> str:
    return datetime.now(timezone.utc).isoformat()


def _text(value) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _nonneg_int(value, default: int = 0) -> int:
    if value is None or value == "":
        return default
    try:
        number = int(round(float(str(value).strip().replace(",", "."))))
    except (TypeError, ValueError) as exc:
        raise ValueError("payment_term_days must be a number") from exc
    if number < 0:
        raise ValueError("payment_term_days must be >= 0")
    return number


def _normalize_interactions(items: list | None) -> list[dict]:
    normalized = []
    for item in items or []:
        if hasattr(item, "model_dump"):
            item = item.model_dump()
        if not isinstance(item, dict):
            continue
        content = _text(item.get("content"))
        if not content:
            continue
        created_at = item.get("created_at")
        if created_at is None:
            created_at = _now()
        elif hasattr(created_at, "isoformat"):
            if getattr(created_at, "tzinfo", None) is not None:
                created_at = created_at.astimezone(timezone.utc)
            created_at = created_at.isoformat()
        else:
            created_at = str(created_at)
        interaction_id = item.get("id") or uuid4()
        normalized.append(
            {
                "id": str(interaction_id),
                "content": content,
                "created_at": created_at,
            }
        )
    return _sort_interactions(normalized)


def _interaction_timestamp(item: dict) -> float:
    raw = item.get("created_at")
    if raw is None or raw == "":
        return float("-inf")
    if hasattr(raw, "timestamp"):
        return raw.timestamp()
    text = str(raw).replace("Z", "+00:00")
    try:
        return datetime.fromisoformat(text).timestamp()
    except ValueError:
        return float("-inf")


def _sort_interactions(items: list | None) -> list[dict]:
    rows = [item for item in (items or []) if isinstance(item, dict)]
    return sorted(rows, key=_interaction_timestamp, reverse=True)


def _format_customer(row: dict) -> dict:
    row["interaction_history"] = _sort_interactions(row.get("interaction_history"))
    row["payment_term_days"] = _nonneg_int(row.get("payment_term_days"))
    return row


def _to_row(
    payload: dict,
    clerk_id: str,
    *,
    include_history: bool,
    default_payment_term_days: int,
) -> dict:
    row = {field: _text(payload.get(field)) for field in TEXT_FIELDS}
    row["payment_term_days"] = _nonneg_int(
        payload.get("payment_term_days"),
        default_payment_term_days,
    )
    row["clerk_id"] = clerk_id
    if include_history:
        row["interaction_history"] = _normalize_interactions(
            payload.get("interaction_history")
        )
    return row


def _owned_query(customer_id: str, clerk_id: str):
    return (
        supabase.table(TABLE)
        .select("*")
        .eq("id", customer_id)
        .eq("clerk_id", clerk_id)
        .limit(1)
    )


def _fold_header(value: str) -> str:
    text = value.strip().lower().replace("đ", "d")
    normalized = unicodedata.normalize("NFD", text)
    stripped = "".join(ch for ch in normalized if unicodedata.category(ch) != "Mn")
    return re.sub(r"[^a-z0-9]+", " ", stripped).strip()


def _cell_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return ""
    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))
        return str(value).strip()
    if isinstance(value, int):
        return str(value)
    return str(value).strip()


def _map_headers(headers: tuple) -> dict[int, str]:
    mapping = {}
    for index, header in enumerate(headers):
        folded = _fold_header(_cell_text(header))
        if not folded:
            continue
        for field, aliases in HEADER_ALIASES.items():
            if folded in aliases:
                mapping[index] = field
                break
    return mapping


def _parse_local_datetime(date_text: str, time_text: str | None) -> datetime | None:
    if "-" in date_text and date_text[0:4].isdigit():
        year_text, month_text, day_text = date_text.split("-", 2)
    else:
        day_text, month_text, year_text = re.split(r"[/-]", date_text)
    hour, minute = (0, 0)
    if time_text:
        hour_text, minute_text = time_text.split(":", 1)
        hour, minute = int(hour_text), int(minute_text)
    try:
        return datetime(
            int(year_text),
            int(month_text),
            int(day_text),
            hour,
            minute,
            tzinfo=_LOCAL_TZ,
        )
    except ValueError:
        return None


def _split_interaction_line(line: str) -> tuple[datetime | None, str]:
    match = _INTERACTION_LINE.match(line.strip())
    if not match:
        return None, line.strip()
    date_text, time_text, rest = match.group(1), match.group(2), match.group(3).strip()
    content = re.sub(r"^[|\-–—:]\s*", "", rest).strip()
    if not content:
        return None, line.strip()
    created_at = _parse_local_datetime(date_text, time_text)
    if created_at is None:
        return None, line.strip()
    return created_at, content


def _parse_interaction_cell(value) -> list[dict]:
    text = _cell_text(value)
    items = []
    for line in re.split(r"\r?\n+", text):
        content_line = line.strip()
        if not content_line:
            continue
        created_at, content = _split_interaction_line(content_line)
        if not content:
            continue
        items.append({"content": content, "created_at": created_at})
    return items


def _row_payload(values: tuple, mapping: dict[int, str]) -> dict:
    payload = {field: "" for field in TEXT_FIELDS}
    payload["payment_term_days"] = None
    payload["interaction_history"] = []
    for index, field in mapping.items():
        if index >= len(values):
            continue
        if field == "interaction_history":
            payload["interaction_history"] = _parse_interaction_cell(values[index])
            continue
        text = _cell_text(values[index])
        if field == "payment_term_days" and text == "":
            continue
        payload[field] = text
    return payload


def _is_empty_payload(payload: dict) -> bool:
    return not any(_text(payload.get(field)) for field in TEXT_FIELDS)


def _identity_key(company_name, tax_code) -> tuple[str, str] | None:
    company = re.sub(r"\s+", " ", _text(company_name)).casefold()
    tax = re.sub(r"[\s.\-]", "", _text(tax_code)).casefold()
    if not company and not tax:
        return None
    return (company, tax)


def _existing_customer_keys(clerk_id: str) -> set[tuple[str, str]]:
    keys: set[tuple[str, str]] = set()
    page_size = 1000
    start = 0
    while True:
        result = (
            supabase.table(TABLE)
            .select("company_name, tax_code")
            .eq("clerk_id", clerk_id)
            .range(start, start + page_size - 1)
            .execute()
        )
        rows = result.data or []
        for row in rows:
            key = _identity_key(row.get("company_name"), row.get("tax_code"))
            if key:
                keys.add(key)
        if len(rows) < page_size:
            break
        start += page_size
    return keys


def create_customers(clerk_id: str, payloads: list[dict]) -> list[dict]:
    if not payloads:
        return []
    ensure_user(clerk_id)
    default_days = get_settings(clerk_id)["default_payment_term_days"]
    saved = []
    for start in range(0, len(payloads), INSERT_CHUNK_SIZE):
        chunk = [
            _to_row(
                item,
                clerk_id,
                include_history=True,
                default_payment_term_days=default_days,
            )
            for item in payloads[start : start + INSERT_CHUNK_SIZE]
        ]
        result = supabase.table(TABLE).insert(chunk).execute()
        saved.extend(_format_customer(row) for row in (result.data or []))
    return saved


def import_customers_from_xlsx(clerk_id: str, content: bytes) -> dict:
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise CustomerImportError("File Excel không hợp lệ.") from exc

    try:
        sheet = workbook.active
        rows = sheet.iter_rows(values_only=True)
        try:
            headers = next(rows)
        except StopIteration as exc:
            raise CustomerImportError("File Excel trống.") from exc

        mapping = _map_headers(headers or ())
        if not mapping:
            raise CustomerImportError(
                "Không nhận diện được cột. Hãy dùng file mẫu (Tên công ty, Điện thoại, Người liên hệ, Địa chỉ, Mã số thuế, Lịch sử tương tác)."
            )

        payloads = []
        skipped = 0
        duplicates = 0
        seen = _existing_customer_keys(clerk_id)
        for values in rows:
            values = values or ()
            if all(value is None or str(value).strip() == "" for value in values):
                skipped += 1
                continue
            payload = _row_payload(values, mapping)
            if _is_empty_payload(payload):
                skipped += 1
                continue
            key = _identity_key(payload.get("company_name"), payload.get("tax_code"))
            if key and key in seen:
                duplicates += 1
                continue
            if key:
                seen.add(key)
            payloads.append(payload)
    finally:
        workbook.close()

    saved = create_customers(clerk_id, payloads)
    return {
        "imported": len(saved),
        "skipped": skipped,
        "duplicates": duplicates,
        "customers": saved,
    }


def build_customer_template() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Khach hang"
    sheet.append(list(TEMPLATE_HEADERS))
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def create_customer(clerk_id: str, payload: dict) -> dict:
    ensure_user(clerk_id)
    default_days = get_settings(clerk_id)["default_payment_term_days"]
    result = (
        supabase.table(TABLE)
        .insert(
            _to_row(
                payload,
                clerk_id,
                include_history=True,
                default_payment_term_days=default_days,
            )
        )
        .execute()
    )
    if not result.data:
        raise RuntimeError("Failed to create customer")
    return _format_customer(result.data[0])


def list_customers(clerk_id: str) -> list[dict]:
    result = (
        supabase.table(TABLE)
        .select("*")
        .eq("clerk_id", clerk_id)
        .order("created_at", desc=True)
        .execute()
    )
    return [_format_customer(row) for row in (result.data or [])]


def get_customer(customer_id: str, clerk_id: str) -> dict:
    result = _owned_query(customer_id, clerk_id).execute()
    data = result.data or []
    if not data:
        raise CustomerNotFoundError("Customer not found")
    return _format_customer(data[0])


def update_customer(customer_id: str, clerk_id: str, payload: dict) -> dict:
    get_customer(customer_id, clerk_id)

    updates = {}
    for field in TEXT_FIELDS:
        if field in payload and payload[field] is not None:
            updates[field] = _text(payload[field])
    if "payment_term_days" in payload and payload["payment_term_days"] is not None:
        updates["payment_term_days"] = _nonneg_int(payload["payment_term_days"])
    if "interaction_history" in payload and payload["interaction_history"] is not None:
        updates["interaction_history"] = _normalize_interactions(
            payload["interaction_history"]
        )
    if not updates:
        return get_customer(customer_id, clerk_id)

    updates["updated_at"] = _now()
    result = (
        supabase.table(TABLE)
        .update(updates)
        .eq("id", customer_id)
        .eq("clerk_id", clerk_id)
        .execute()
    )
    if not result.data:
        raise CustomerNotFoundError("Customer not found")
    return _format_customer(result.data[0])


def delete_customer(customer_id: str, clerk_id: str) -> dict:
    existing = get_customer(customer_id, clerk_id)
    # Orders and quotes reference this customer. Remove them first so stock
    # reserved by a delivered order is returned before the order disappears.
    from src.services.orders import delete_order

    order_rows = (
        supabase.table("orders")
        .select("id")
        .eq("customer_id", customer_id)
        .eq("clerk_id", clerk_id)
        .execute()
    )
    for row in order_rows.data or []:
        delete_order(row["id"], clerk_id)
    (
        supabase.table("quotes")
        .delete()
        .eq("customer_id", customer_id)
        .eq("clerk_id", clerk_id)
        .execute()
    )
    supabase.table(TABLE).delete().eq("id", customer_id).eq(
        "clerk_id", clerk_id
    ).execute()
    return existing
